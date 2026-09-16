"""
build.py — Assembles docs/sections/*.md into a formatted .docx

Usage:
    python docs/build.py
    python docs/build.py --output my_output.docx

Output: GSP_PULL_Customer_Forecast_Process_Documentation.docx (in project root)

Requirements: pip install python-docx markdown
"""

import argparse
import glob
import os
import re

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


SECTIONS_DIR = os.path.join(os.path.dirname(__file__), "sections")
DEFAULT_OUTPUT = os.path.join(
    os.path.dirname(__file__), "..", "GSP_PULL_Customer_Forecast_Process_Documentation.docx"
)


def set_heading_style(para, level):
    """Apply heading style."""
    style_map = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}
    para.style = style_map.get(level, "Heading 1")


def add_table_from_md(doc, lines, start):
    """Parse a markdown table starting at `start` and add it to the doc. Returns next line index."""
    table_lines = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        table_lines.append(lines[i].strip())
        i += 1

    if len(table_lines) < 2:
        return i

    # Parse header row
    header_cells = [c.strip() for c in table_lines[0].split("|") if c.strip()]
    num_cols = len(header_cells)

    # Skip separator row (index 1), collect data rows
    data_rows = []
    for row_line in table_lines[2:]:
        cells = [c.strip() for c in row_line.split("|") if c.strip() != ""]
        # Pad or trim to num_cols
        while len(cells) < num_cols:
            cells.append("")
        data_rows.append(cells[:num_cols])

    table = doc.add_table(rows=1 + len(data_rows), cols=num_cols)
    table.style = "Table Grid"

    # Header row — bold + light blue background
    hdr_row = table.rows[0]
    for j, text in enumerate(header_cells):
        cell = hdr_row.cells[j]
        cell.text = text
        run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.paragraphs[0].add_run(text)
        run.bold = True
        # Light blue shading
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "D6E4F0")
        tc_pr.append(shd)

    # Data rows
    for r, row_data in enumerate(data_rows):
        row = table.rows[r + 1]
        for j, text in enumerate(row_data):
            row.cells[j].text = text

    doc.add_paragraph()  # spacing after table
    return i


def inline_format(para, text):
    """Add a run to para, handling **bold** and `code` inline markers."""
    # Split on bold and code markers
    pattern = r"(\*\*[^*]+\*\*|`[^`]+`)"
    parts = re.split(pattern, text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = para.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        else:
            para.add_run(part)


def process_section(doc, md_path):
    """Parse a markdown file and add its content to the docx."""
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")
        stripped = line.strip()

        # Headings
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            set_heading_style(p, 3)
            p.clear()
            inline_format(p, stripped[4:])
            i += 1

        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            set_heading_style(p, 2)
            p.clear()
            inline_format(p, stripped[3:])
            i += 1

        elif stripped.startswith("# "):
            p = doc.add_paragraph()
            set_heading_style(p, 1)
            p.clear()
            inline_format(p, stripped[2:])
            i += 1

        # Table
        elif stripped.startswith("|"):
            i = add_table_from_md(doc, lines, i)

        # Code block
        elif stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            i += 1  # skip closing ```
            p = doc.add_paragraph()
            p.style = "No Spacing"
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            doc.add_paragraph()

        # Blockquote
        elif stripped.startswith("> "):
            content = stripped[2:]
            p = doc.add_paragraph()
            p.style = "Quote" if "Quote" in [s.name for s in doc.styles] else "Normal"
            inline_format(p, content)
            i += 1

        # Bullet list
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            inline_format(p, stripped[2:])
            i += 1

        # Numbered list
        elif re.match(r"^\d+\.\s", stripped):
            content = re.sub(r"^\d+\.\s", "", stripped)
            p = doc.add_paragraph(style="List Number")
            inline_format(p, content)
            i += 1

        # Horizontal rule
        elif stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            i += 1

        # Empty line
        elif stripped == "":
            i += 1

        # Normal paragraph
        else:
            p = doc.add_paragraph()
            inline_format(p, stripped)
            i += 1


def build(output_path):
    doc = Document()

    # Title page paragraph
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("GSP — Demand Forecast Snapshot\nMaster Process Documentation")
    run.bold = True
    run.font.size = Pt(20)
    doc.add_paragraph()

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Snapshot Logic · Table Lifecycle · Batch Schedule · Exception Handling"
    ).italic = True
    doc.add_paragraph()

    # Collect and sort section files
    section_files = sorted(glob.glob(os.path.join(SECTIONS_DIR, "*.md")))
    if not section_files:
        print(f"No .md files found in {SECTIONS_DIR}")
        return

    for md_file in section_files:
        process_section(doc, md_file)
        doc.add_paragraph()  # space between sections

    output_path = os.path.abspath(output_path)
    doc.save(output_path)
    print(f"✅  Built: {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build GSP forecast doc from markdown sections.")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Output .docx path")
    args = parser.parse_args()
    build(args.output)
